"""Read-only problem inputs; all generated artifacts go to a separate output dir."""
from dataclasses import asdict
from pathlib import Path
import csv
import hashlib
import json
import numpy as np
import openpyxl
from .model import Field

ROOT=Path(__file__).resolve().parent.parent
PROBLEM=ROOT/'problem'
HEADERS=['吸收塔x坐标 (m)','吸收塔y坐标 (m)','定日镜序号','定日镜宽度 (m)',
         '定日镜高度 (m)','定日镜x坐标 (m)','定日镜y坐标 (m)','定日镜z坐标 (m)']
METRIC_COLUMNS=['optical_eff','cosine_eff','shadow_block_eff','atmospheric_eff',
                'trunc_eff','power_MW','unit_power_kW_m2','unshadowed_eff',
                'sequential_blocking_loss','isolated_trunc_eff']


def load_attachment():
    wb=openpyxl.load_workbook(PROBLEM/'附件.xlsx',read_only=True,data_only=True)
    try:
        xy=np.array([row[:2] for row in wb.active.iter_rows(min_row=2,values_only=True)
                     if row[0] is not None],float)
    finally:
        wb.close()
    return Field(np.c_[xy,np.full(len(xy),4.)],6.,6.,np.zeros(2),'q1')


def safe_output(path):
    path=Path(path).resolve()
    if path==PROBLEM.resolve() or PROBLEM.resolve() in path.parents:
        raise ValueError('Writing inside problem/ is prohibited')
    path.parent.mkdir(parents=True,exist_ok=True)
    return path


def save_field(field,path):
    path=safe_output(path)
    np.savez_compressed(path,centers=field.centers,widths=field.widths,
                        heights=field.heights,tower=field.tower,name=field.name)


def load_field(path):
    with np.load(path,allow_pickle=False) as f:
        return Field(f['centers'],f['widths'],f['heights'],f['tower'],str(f['name']))


def write_design(field,path):
    path=safe_output(path)
    # Repository solver output, not editing the supplied workbook templates.
    wb=openpyxl.Workbook()
    ws=wb.active;ws.title='定日镜参数'
    ws.append(HEADERS)
    for i,(c,w,h) in enumerate(zip(field.centers,field.widths,field.heights),1):
        ws.append([float(field.tower[0]),float(field.tower[1]),i,float(w),float(h),*map(float,c)])
    from openpyxl.styles import Font,PatternFill,Alignment
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='174A5B')
        cell.alignment=Alignment(horizontal='center')
    for column in 'ABCDEFGH':
        ws.column_dimensions[column].width=25
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column!=3:cell.number_format='0.000000'
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(path);wb.close()


def write_json(value,path):
    def default(o):
        if isinstance(o,np.ndarray):return o.tolist()
        if isinstance(o,np.generic):return o.item()
        if hasattr(o,'__dataclass_fields__'):return asdict(o)
        raise TypeError(type(o).__name__)
    safe_output(path).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=default,allow_nan=False)+'\n')


def input_hashes():
    return {str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(PROBLEM.iterdir()) if p.is_file()}


def write_metrics(evaluations,path):
    with safe_output(path).open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['problem','month',*METRIC_COLUMNS])
        for name,result in evaluations.items():
            for month,row in enumerate(result.monthly(),1):
                writer.writerow([name,month,*row])


def write_instantaneous(evaluations,path):
    from .model import time_grid
    labels,_,_=time_grid()
    with safe_output(path).open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(['problem','month','local_time',*METRIC_COLUMNS])
        for name,result in evaluations.items():
            for (month,t),row in zip(labels,result.rows()):
                writer.writerow([name,month,t,*row])
